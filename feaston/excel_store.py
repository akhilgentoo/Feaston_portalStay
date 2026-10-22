"""Tiny Excel (.xlsx) table storage: first row is the header, each following row is a record."""

from __future__ import annotations

import json
from pathlib import Path

from openpyxl import Workbook, load_workbook


def _to_cell(value):
    if isinstance(value, (dict, list, tuple)):
        return json.dumps(value, ensure_ascii=True)
    return value


def read_rows(path: Path) -> list[dict]:
    if not path.exists():
        return []
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        rows = workbook.active.iter_rows(values_only=True)
        headers = next(rows, None)
        if not headers:
            return []
        keys = [str(h).strip() if h is not None else "" for h in headers]
        records: list[dict] = []
        for row in rows:
            if row is None or all(cell is None for cell in row):
                continue
            records.append({key: ("" if cell is None else cell) for key, cell in zip(keys, row) if key})
        return records
    finally:
        workbook.close()


def write_rows(path: Path, headers: list[str], rows: list[dict], sheet_title: str = "Sheet1") -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = sheet_title
    sheet.append(headers)
    for row in rows:
        sheet.append([_to_cell(row.get(header, "")) for header in headers])
    workbook.save(path)


def append_row(path: Path, headers: list[str], row: dict, sheet_title: str = "Sheet1") -> None:
    if not path.exists():
        write_rows(path, headers, [row], sheet_title)
        return
    workbook = load_workbook(path)
    sheet = workbook.active
    sheet.append([_to_cell(row.get(header, "")) for header in headers])
    workbook.save(path)
